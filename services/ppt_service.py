from pptx.util import Cm
from pptx.util import Pt


def create_table(shapes, body_data):
    rows, cols = len(body_data) + 1, len(body_data[0])
    left, top, width, height = Cm(1), Cm(4), Cm(rows * 24 / 14), Cm(cols * 11 / 6)
    return shapes.add_table(rows, cols, left, top, width, height).table


def write_table(ppt_table, headings, body_data):
    for heading_index, heading_item in enumerate(headings, start=0):
        cell = ppt_table.cell(0, heading_index)
        cell.text = heading_item
        set_cell_font_size(cell, 14)

    for body_index, body_item in enumerate(body_data, start=1):
        for col_index, text in enumerate(body_item, start=0):
            cell = ppt_table.cell(body_index, col_index)
            cell.text = str(text)
            set_cell_font_size(cell, 12)


def set_cell_font_size(cell, size):
    for paragraph in cell.text_frame.paragraphs:
        for paragraph_run in paragraph.runs:
            paragraph_run.font.size = Pt(size)
